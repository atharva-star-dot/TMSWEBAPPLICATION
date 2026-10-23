"""
Automated checks for https://tmswebapplication.vercel.app
- Reads test cases/test data from the Excel file below
- Writes results to results.csv AND to a copy of the Excel file (Actual Result, Status, Remarks)

Setup (once):
    pip install playwright openpyxl
    playwright install chromium
Run:  python test.py          (set HEADED = True below to watch the browser)
"""
import os, re, csv
import openpyxl
from openpyxl.styles import PatternFill
from playwright.sync_api import sync_playwright

# ---------------- CONFIG (edit here) ----------------
BASE   = "https://tmswebapplication.vercel.app"
EMAIL  = "atharvavadane12@gmail.com"      # valid test account
PWD    = "Atharva@12"
HEADED = False                             # True = show browser window
XLSX_PATH = r"E:\OneMoreTime\TMS\TMSWEBAPPLICATION\testing\TMS_WebApp_Test_Cases.xlsx"
OUT_XLSX  = XLSX_PATH.replace(".xlsx", "_results.xlsx")   # original file stays untouched
OUT_CSV   = os.path.join(os.path.dirname(XLSX_PATH), "results.csv")
# ----------------------------------------------------

# ---------- load test cases from Excel ----------
wb = openpyxl.load_workbook(XLSX_PATH)
ws_all = wb["All Test Cases"]
hdr = [c.value for c in ws_all[1]]
COL = {h: i for i, h in enumerate(hdr)}
TESTS = {}   # TC ID -> {"email":..., "password":..., "raw": Test Data}
for r in ws_all.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    raw = str(r[COL["Test Data"]] or "")
    def pick(label):
        m = re.search(label + r":\s*(.*?)(?:\s*\|\s*\w+:|$)", raw)
        if not m: return None
        v = m.group(1).strip()
        return "" if v.lower() == "(blank)" else v
    TESTS[r[0]] = {"raw": raw, "email": pick("Email"), "password": pick("Password")}

def td(tid, field, default=None):
    """Value from the sheet's Test Data column for this TC, else default."""
    v = TESTS.get(tid, {}).get(field)
    return default if v is None else v

print(f"Loaded {len(TESTS)} test cases from Excel")

rows = []

def record(tid, ok, actual, remark=""):
    rows.append([tid, actual, "Pass" if ok else "Fail", "" if ok else (remark or "See Actual Result")])
    print(tid, "Pass" if ok else "Fail", "-", actual)

def blocked(tid, err):
    rows.append([tid, f"Test could not run: {str(err)[:150]}", "Blocked", "Check selector/network; rerun or test manually"])
    print(tid, "Blocked", err)

def login_page(ctx):
    p = ctx.new_page(); p.goto(BASE + "/login"); return p

def submit(p, email, pw, press_enter=False):
    p.fill('input[type="email"]', email); p.fill('input[type="password"]', pw)
    if press_enter: p.press('input[type="password"]', "Enter")
    else: p.get_by_role("button", name=re.compile("Sign In", re.I)).click()

def on_dashboard(p):
    try: p.wait_for_url("**/dashboard**", timeout=15000); return True
    except Exception: return False

with sync_playwright() as pw:
    br = pw.chromium.launch(headless=not HEADED)
    def fresh(): return br.new_context()
    def run(tid, fn):
        try: fn()
        except Exception as e: blocked(tid, e)

    def t001():
        p = login_page(fresh()); submit(p, EMAIL, PWD); ok = on_dashboard(p)
        record("TC_LOGIN_001", ok, f"After valid login URL is {p.url}", "Valid login did not reach /dashboard (Critical)")
    def bad_login(tid, label, fallback_email, fallback_pw):
        # email/password come from the Excel "Test Data" column; the valid account is used when the sheet has a placeholder
        email = td(tid, "email", fallback_email)
        pw_   = td(tid, "password", fallback_pw)
        if email == "qa.admin@example.com": email = EMAIL      # sheet placeholder -> real account
        p = login_page(fresh()); submit(p, email, pw_); p.wait_for_timeout(4000)
        err = p.locator("text=/invalid|incorrect|wrong|error|failed/i").count() > 0
        ok = "/login" in p.url and err
        record(tid, ok, f"{label}: URL {p.url}; error message shown: {err}", f"{label} not handled as expected")
    def t004():
        p = login_page(fresh()); p.get_by_role("button", name=re.compile("Sign In", re.I)).click()
        miss = p.locator('input[type="email"]').evaluate("e=>e.validity.valueMissing")
        record("TC_LOGIN_004", miss and "/login" in p.url, f"Blank submit: email required flag={miss}, URL {p.url}")
    def t005():
        p = login_page(fresh()); p.fill('input[type="email"]', EMAIL)
        p.get_by_role("button", name=re.compile("Sign In", re.I)).click()
        miss = p.locator('input[type="password"]').evaluate("e=>e.validity.valueMissing")
        record("TC_LOGIN_005", miss and "/login" in p.url, f"Password blank: required flag={miss}, URL {p.url}")
    def payload(tid, email, label):
        p = login_page(fresh()); fired = []
        p.on("dialog", lambda d: (fired.append(d.message), d.dismiss()))
        submit(p, email, td(tid, "password", "x")); p.wait_for_timeout(4000)
        ok = not fired and "/dashboard" not in p.url
        record(tid, ok, f"{label}: dialog fired={bool(fired)}, URL {p.url}", f"{label} payload misbehaved (Critical)")
    def t014():
        p = login_page(fresh()); submit(p, EMAIL, PWD, press_enter=True)
        record("TC_LOGIN_014", on_dashboard(p), f"Enter key login URL {p.url}")
    def t016():
        p = login_page(fresh()); p.get_by_role("link", name=re.compile("Create an account", re.I)).click()
        p.wait_for_timeout(1500); record("TC_LOGIN_016", p.url.endswith("/signup"), f"URL {p.url}")
    def t017():
        p = fresh().new_page(); p.goto(BASE + "/dashboard"); p.wait_for_timeout(2000)
        record("TC_LOGIN_017", "/login" in p.url, f"Logged-out /dashboard ended at {p.url}", "Protected route not redirecting (Critical)")
    def t021():
        p = login_page(fresh()); n = p.get_by_text(re.compile("forgot", re.I)).count()
        record("TC_LOGIN_021", n > 0, f"Forgot-password links found: {n}", "No recovery path (Minor)")
    def t_sec001():
        r = fresh().request.get(BASE.replace("https://", "http://"), max_redirects=0)
        loc = r.headers.get("location", "")
        record("TC_SEC_001", r.status in (301, 302, 307, 308) and loc.startswith("https://"), f"HTTP request status {r.status}, Location {loc}")
    def t_sec002():
        r = fresh().new_page().goto(BASE); h = {k.lower(): v for k, v in r.headers.items()}
        need = ["strict-transport-security", "x-content-type-options", "referrer-policy"]
        miss = [x for x in need if x not in h]
        if "x-frame-options" not in h and "frame-ancestors" not in h.get("content-security-policy", ""): miss.append("x-frame-options/CSP frame-ancestors")
        record("TC_SEC_002", not miss, "Missing headers: " + (", ".join(miss) or "none"), "Add security headers in next.config (Major)")
    def t_404():
        p = fresh().new_page(); r = p.goto(BASE + "/xyz-404-check")
        record("TC_DEPLOY_011", r.status == 404, f"Status {r.status}; page text: {p.inner_text('body')[:80]!r}", "Check page is branded manually")
    def exists(tid, path):
        r = fresh().request.get(BASE + path); record(tid, r.status == 200, f"{path} returned {r.status}", f"{path} missing (Minor)")

    run("TC_LOGIN_001", t001)
    run("TC_LOGIN_002", lambda: bad_login("TC_LOGIN_002", "Wrong password", EMAIL, "Wrong@999"))
    run("TC_LOGIN_003", lambda: bad_login("TC_LOGIN_003", "Unregistered email", "nouser_98231@example.com", "Test@12345"))
    run("TC_LOGIN_004", t004); run("TC_LOGIN_005", t005)
    run("TC_LOGIN_010", lambda: payload("TC_LOGIN_010", td("TC_LOGIN_010", "email", "' OR '1'='1' --"), "SQLi"))
    run("TC_LOGIN_011", lambda: payload("TC_LOGIN_011", "<script>alert(1)</script>@x.com", "XSS"))
    run("TC_LOGIN_014", t014); run("TC_LOGIN_016", t016); run("TC_LOGIN_017", t017); run("TC_LOGIN_021", t021)
    run("TC_SEC_001", t_sec001); run("TC_SEC_002", t_sec002); run("TC_DEPLOY_011", t_404)
    run("TC_DEPLOY_014", lambda: exists("TC_DEPLOY_014", "/robots.txt"))
    run("TC_DEPLOY_015", lambda: exists("TC_DEPLOY_015", "/sitemap.xml"))
    br.close()

# ---------- write results.csv ----------
with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f, quoting=csv.QUOTE_ALL); w.writerow(["TC ID", "Actual Result", "Status", "Remarks"]); w.writerows(rows)

# ---------- write results back into a copy of the Excel (all sheets, matched by TC ID) ----------
res = {r[0]: r for r in rows}
FILL = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFD966"}
for ws in wb.worksheets:
    h = [c.value for c in ws[1]]
    if "TC ID" not in h or "Actual Result" not in h:
        continue
    ci, ca, cs, cr = (h.index("TC ID"), h.index("Actual Result"),
                      next(i for i, x in enumerate(h) if str(x).startswith("Status")), h.index("Remarks"))
    for row in ws.iter_rows(min_row=2):
        r = res.get(row[ci].value)
        if r:
            row[ca].value, row[cs].value, row[cr].value = r[1], r[2], r[3]
            row[cs].fill = PatternFill("solid", fgColor=FILL[r[2]])
wb.save(OUT_XLSX)
print(f"\nWrote {len(rows)} results -> {OUT_CSV}\nUpdated copy -> {OUT_XLSX}")